"""文档解析：把 docx 变成可直接检查的结构。

只做一件事——把「一堆段落」变成「章节树 + 元信息」。
不做任何判断，判断留给 checker.py。
"""

from __future__ import annotations

import re
from typing import IO

from docx import Document

from core.schemas import DocProfile, Paragraph, Section

# 章节标题形如：3.1 引言 / 3.5.2 主要数据表结构 / 2 系统设计
SECTION_RE = re.compile(r"^(\d+(?:\.\d+){0,2})\s*[、.．]?\s*(\S.{0,38})$")


def _as_heading(text: str) -> tuple[str, str] | None:
    """判断一个段落是不是章节标题。

    三重限制，避免把正文误判成标题：
    1. 长度不超过 40 字
    2. 不以句末标点结尾（标题不会是完整句子）
    3. 必须以章节编号开头，且编号后还有标题文字
    """
    if len(text) > 40 or text.endswith(("。", "；", "，", "：", "!", "？", ".", ";")):
        return None
    m = SECTION_RE.match(text)
    if not m:
        return None
    number, title = m.group(1), m.group(2).strip()
    if len(title) < 2:
        return None
    # 标题里出现这些，多半是正文里的序号列举
    if title.startswith(("、", "）", ")", "．")):
        return None
    return number, title


def _split_sections(paras: list[Paragraph]) -> list[Section]:
    marks: list[tuple[int, str, str]] = []
    for i, p in enumerate(paras):
        got = _as_heading(p.text)
        if got:
            marks.append((i, got[0], got[1]))

    sections: list[Section] = []
    for j, (idx, number, title) in enumerate(marks):
        end = marks[j + 1][0] if j + 1 < len(marks) else len(paras)
        body = [paras[k].text for k in range(idx + 1, end)]
        sections.append(Section(number=number, title=title, body=body))
    return sections


def parse_docx(source: str | IO[bytes]) -> DocProfile:
    """解析 docx。source 可以是路径，也可以是上传得到的文件对象。"""
    doc = Document(source)

    paras: list[Paragraph] = []
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text:
            paras.append(Paragraph(index=i, text=text, style=p.style.name))

    name = getattr(source, "name", None) or str(source)
    return DocProfile(
        filename=str(name).split("/")[-1].split("\\")[-1],
        paragraphs=paras,
        sections=_split_sections(paras),
        image_count=len(doc.inline_shapes),
        table_count=len(doc.tables),
    )
